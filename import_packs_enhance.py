"""CA FirmHub — Excel import with Compliance Packs.
Designed & Built by Harsh Pala.
Overrides clients import template + excel handler.
"""
from datetime import date
from flask import request, jsonify, g, Response


ALLOWED_PACKS = ("GST_MONTHLY", "GST_QRMP", "TDS", "ITR", "ROC", "PF_ESI")


def register_import_packs_enhance(app, helpers):
    qry = helpers["qry"]
    execute = helpers["execute"]
    login_required = helpers["login_required"]
    qry_id = helpers.get("qry_id")
    log_action = helpers.get("log_action", lambda *a, **k: None)

    def _col_names(table):
        try:
            rows = qry(
                "SELECT column_name FROM information_schema.columns "
                "WHERE table_schema='public' AND table_name=%s", (table,),
            ) or []
            return {str(r.get("column_name") or "").lower() for r in rows}
        except Exception:
            return set()

    def _ensure():
        try:
            execute("ALTER TABLE clients ADD COLUMN IF NOT EXISTS compliance_packs TEXT DEFAULT ''")
        except Exception:
            cols = _col_names("clients")
            if cols and "compliance_packs" not in cols:
                try:
                    execute("ALTER TABLE clients ADD COLUMN compliance_packs TEXT DEFAULT ''")
                except Exception as e:
                    print("[import-packs] col:", e)
        for sql in (
            "ALTER TABLE compliance_calendar ADD COLUMN IF NOT EXISTS client_id INTEGER",
            "ALTER TABLE compliance_calendar ADD COLUMN IF NOT EXISTS source TEXT DEFAULT 'manual'",
        ):
            try:
                execute(sql)
            except Exception:
                pass

    def _parse_fy(fy):
        parts = (fy or "").strip().replace("/", "-").split("-")
        sy = int(parts[0])
        ey = int(parts[1]) if len(parts) > 1 and parts[1].isdigit() else sy + 1
        if ey < 100:
            ey = sy // 100 * 100 + ey if ey > (sy % 100) else (sy // 100 + 1) * 100 + ey
        if ey == sy:
            ey = sy + 1
        return sy, ey

    def _parse_packs_cell(text):
        if not text:
            return []
        raw = str(text).upper().replace(";", ",").replace("|", ",").replace("/", ",")
        aliases = {
            "GST": "GST_MONTHLY", "GSTMONTHLY": "GST_MONTHLY", "GST_MONTHLY": "GST_MONTHLY",
            "QRMP": "GST_QRMP", "GSTQRMP": "GST_QRMP", "GST_QRMP": "GST_QRMP",
            "TDS": "TDS", "ITR": "ITR", "INCOME_TAX": "ITR", "INCOMETAX": "ITR",
            "ROC": "ROC", "PF": "PF_ESI", "ESI": "PF_ESI", "PF_ESI": "PF_ESI", "PFESI": "PF_ESI",
        }
        out = []
        for part in raw.split(","):
            p = part.strip().replace(" ", "_")
            p2 = aliases.get(p, p)
            if p2 in ALLOWED_PACKS and p2 not in out:
                out.append(p2)
        return out

    def _parse_engagement_types(cell):
        if cell is None:
            return []
        text = str(cell).strip()
        if not text or text.lower() == "nan":
            return []
        low = text.lower()
        keywords = [
            ("statutory", "Statutory Audit"), ("tax audit", "Tax Audit"),
            ("internal", "Internal Audit"), ("gst", "GST Audit"),
            ("due diligence", "Due Diligence"), ("certification", "Certification"),
        ]
        found = []
        for kw, can in keywords:
            if kw in low and can not in found:
                found.append(can)
        return found if found else [f"Other:{text}"]

    def _client_deadlines(packs, name, sy, ey):
        packs = set((packs or "").upper().replace(" ", "").split(","))
        packs.discard("")
        items, name = [], name or "Client"
        if "GST_MONTHLY" in packs:
            for y, m in [(sy, m) for m in range(4, 13)] + [(ey, m) for m in range(1, 4)]:
                dy, dm = (y + 1, 1) if m == 12 else (y, m + 1)
                mon = date(y, m, 1).strftime("%b %Y")
                items += [
                    (f"{name}: GSTR-1 {mon}", "GST", date(dy, dm, 11).isoformat()),
                    (f"{name}: GSTR-3B {mon}", "GST", date(dy, dm, 20).isoformat()),
                ]
        if "GST_QRMP" in packs:
            items += [
                (f"{name}: GSTR-1 QRMP Q1", "GST", f"{sy}-07-13"),
                (f"{name}: GSTR-1 QRMP Q2", "GST", f"{sy}-10-13"),
                (f"{name}: GSTR-1 QRMP Q3", "GST", f"{ey}-01-13"),
                (f"{name}: GSTR-1 QRMP Q4", "GST", f"{ey}-04-13"),
            ]
        if "TDS" in packs:
            items += [
                (f"{name}: TDS Q1", "TDS", f"{sy}-07-31"), (f"{name}: TDS Q2", "TDS", f"{sy}-10-31"),
                (f"{name}: TDS Q3", "TDS", f"{ey}-01-31"), (f"{name}: TDS Q4", "TDS", f"{ey}-05-31"),
            ]
        if "ITR" in packs:
            items += [(f"{name}: ITR filing", "ITR", f"{ey}-07-31"), (f"{name}: Tax audit", "ITR", f"{ey}-10-31")]
        if "ROC" in packs:
            items += [(f"{name}: AOC-4", "ROC", f"{ey}-10-29"), (f"{name}: MGT-7", "ROC", f"{ey}-11-29")]
        if "PF_ESI" in packs:
            for y, m in [(sy, m) for m in range(4, 13)] + [(ey, m) for m in range(1, 4)]:
                dy, dm = (y + 1, 1) if m == 12 else (y, m + 1)
                mon = date(y, m, 1).strftime("%b %Y")
                items.append((f"{name}: PF/ESI {mon}", "Other", date(dy, dm, 15).isoformat()))
        return items

    def _uid():
        u = g.user or {}
        return u.get("id") or u.get("user_id")

    def _insert_cal(title, cat, due, fy, client_id):
        cols = _col_names("compliance_calendar")
        uid = _uid()
        if "client_id" in cols or not cols:
            try:
                execute(
                    "INSERT INTO compliance_calendar "
                    "(firm_id,title,category,due_date,financial_year,is_recurring,status,created_by_id,client_id,source) "
                    "VALUES (%s,%s,%s,%s,%s,TRUE,'Upcoming',%s,%s,'import_pack')",
                    (g.firm_id, title, cat, due, fy, uid, client_id),
                )
                return
            except Exception:
                pass
        execute(
            "INSERT INTO compliance_calendar "
            "(firm_id,title,category,due_date,financial_year,is_recurring,status,created_by_id) "
            "VALUES (%s,%s,%s,%s,%s,TRUE,'Upcoming',%s)",
            (g.firm_id, title, cat, due, fy, uid),
        )

    def _apply_packs(client_id, client_name, packs_list, fy):
        if not packs_list:
            return 0
        packs_str = ",".join(packs_list)
        try:
            execute("UPDATE clients SET compliance_packs=%s WHERE id=%s AND firm_id=%s",
                    (packs_str, client_id, g.firm_id))
        except Exception as e:
            print("[import-packs] save:", e)
            return 0
        try:
            sy, ey = _parse_fy(fy)
        except Exception:
            return 0
        added = 0
        for title, cat, due in _client_deadlines(packs_str, client_name, sy, ey):
            if qry("SELECT id FROM compliance_calendar WHERE firm_id=%s AND title=%s AND financial_year=%s",
                   (g.firm_id, title, fy), one=True):
                continue
            try:
                _insert_cal(title, cat, due, fy, client_id)
                added += 1
            except Exception as e:
                print("[import-packs] insert:", e)
        return added

    def enhanced_template():
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        import io
        wb = Workbook()
        ws = wb.active
        ws.title = "Clients Import"
        headers = ["Name of Client", "Type of Entity", "PAN", "Contact Person", "Phone",
                   "Email", "Address", "Type of Engagement", "Compliance Packs"]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2563EB")
        ws.append(["ABC Textiles Pvt Ltd", "Company", "ABCDE1234F", "Rajesh Shah", "9876543210",
                   "rajesh@abctextiles.com", "123, MG Road, Mumbai", "Statutory & Tax Audit",
                   "GST_MONTHLY, TDS, ITR, ROC"])
        ws.append(["Sunrise Traders", "Firm", "FGHIJ5678K", "Priya Mehta", "9823456789",
                   "priya@sunrisetraders.com", "45, Market Street, Pune", "GST Audit",
                   "GST_MONTHLY, ITR"])
        ws.append(["Mehta Associates", "LLP", "KLMNO9012P", "Amit Mehta", "9988776655",
                   "amit@mehta.in", "Ahmedabad", "Tax Audit", "ITR, TDS"])
        for i, w in enumerate([26, 14, 14, 18, 14, 26, 28, 26, 32], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        packs_ws = wb.create_sheet("Packs Checklist")
        packs_ws.append(["Tick (Y/N)", "Pack Code", "What it generates on import"])
        for c in packs_ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="003366")
        for row in [
            ("", "GST_MONTHLY", "Monthly GSTR-1 & GSTR-3B for the import FY"),
            ("", "GST_QRMP", "Quarterly GSTR-1 QRMP dates"),
            ("", "TDS", "Quarterly TDS return dates"),
            ("", "ITR", "ITR filing + Tax audit dates"),
            ("", "ROC", "AOC-4 and MGT-7 dates"),
            ("", "PF_ESI", "Monthly PF/ESI dates"),
        ]:
            packs_ws.append(list(row))
        packs_ws.append([])
        packs_ws.append(["How to use", "In column Compliance Packs, write codes separated by commas."])
        packs_ws.append(["Example", "GST_MONTHLY, ITR, ROC"])
        packs_ws.append(["Optional", "Or tick default packs in the Import dialog when Excel cell is blank."])
        packs_ws.column_dimensions["A"].width = 12
        packs_ws.column_dimensions["B"].width = 16
        packs_ws.column_dimensions["C"].width = 55

        notes = wb.create_sheet("Instructions")
        notes.append(["Column", "Notes"])
        notes["A1"].font = notes["B1"].font = Font(bold=True)
        for r in [
            ("Name of Client", "Required."),
            ("Type of Entity", "Company / Firm / LLP / Individual / Trust / HUF / AOP/BOI"),
            ("PAN", "Optional; primary duplicate check."),
            ("Type of Engagement", "e.g. Statutory Audit, Tax Audit, or combined."),
            ("Compliance Packs", "Comma-separated: GST_MONTHLY, GST_QRMP, TDS, ITR, ROC, PF_ESI. "
             "Saved on client; statutory calendar generated for selected FY."),
            ("Financial Year", "Uses Dashboard FY selector at import time."),
        ]:
            notes.append(list(r))
        notes.column_dimensions["A"].width = 28
        notes.column_dimensions["B"].width = 95

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return Response(
            buf.read(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=clients_import_template.xlsx"},
        )

    def enhanced_import():
        try:
            import pandas as pd
        except ImportError:
            return jsonify({"detail": "pandas not installed"}), 500
        _ensure()
        import_fy = (request.form.get("fy") or "").strip()
        if not import_fy:
            today = date.today()
            import_fy = f"{today.year}-{str(today.year+1)[-2:]}" if today.month >= 4 else f"{today.year-1}-{str(today.year)[-2:]}"
        default_packs = _parse_packs_cell(request.form.get("default_packs") or "")
        f = request.files.get("file")
        if not f:
            return jsonify({"detail": "No file uploaded"}), 400
        import io
        try:
            raw = f.read()
            df = pd.read_csv(io.BytesIO(raw)) if f.filename.lower().endswith(".csv") else pd.read_excel(io.BytesIO(raw))
        except Exception as e:
            return jsonify({"detail": f"Could not read file: {e}"}), 400
        df.columns = [str(c).strip() for c in df.columns]
        if "Name of Client" not in df.columns:
            return jsonify({"detail": "Missing required column: Name of Client"}), 400
        packs_col = None
        for cand in ("Compliance Packs", "Compliance packs", "Packs", "compliance_packs", "PACKS"):
            if cand in df.columns:
                packs_col = cand
                break

        def cell(row, col):
            if col is None:
                return None
            try:
                v = row[col]
            except Exception:
                return None
            try:
                if pd.isna(v):
                    return None
            except Exception:
                pass
            v = str(v).strip()
            return v if v and v.lower() != "nan" else None

        ENTITY = {"Company", "Firm", "LLP", "Individual", "Trust", "HUF", "AOP/BOI"}
        existing = qry("SELECT id,name,pan FROM clients WHERE firm_id=%s", (g.firm_id,)) or []
        pan_index = {c["pan"].strip().upper(): c for c in existing if c.get("pan")}
        name_index = {c["name"].strip().lower(): c for c in existing if c.get("name")}
        existing_engs = qry("SELECT client_id,engagement_type,financial_year FROM engagements WHERE firm_id=%s", (g.firm_id,)) or []
        eng_key = lambda cid, et, fy: f"{cid}|{et}|{fy}"
        eng_index = {eng_key(e["client_id"], e["engagement_type"], e["financial_year"]) for e in existing_engs}

        clients_created = clients_skipped = 0
        engagements_created = engagements_skipped = 0
        packs_applied = calendar_added = 0
        row_errors, row_warnings = [], []
        fys_touched = set()

        for i, row in df.iterrows():
            rn = int(i) + 2
            name = cell(row, "Name of Client")
            if not name:
                row_errors.append(f"Row {rn}: Name of Client is required")
                continue
            pan = cell(row, "PAN")
            entity_type = cell(row, "Type of Entity")
            if entity_type and entity_type not in ENTITY:
                row_warnings.append(f"Row {rn}: Unknown entity type '{entity_type}'")
            client_id = None
            if pan and pan.upper() in pan_index:
                client_id = pan_index[pan.upper()]["id"]
                clients_skipped += 1
            elif name.lower() in name_index:
                client_id = name_index[name.lower()]["id"]
                clients_skipped += 1
            else:
                try:
                    args = (g.firm_id, name, pan, cell(row, "Contact Person"), cell(row, "Phone"),
                            cell(row, "Email"), cell(row, "Address"), entity_type, g.user["id"])
                    sql = ("INSERT INTO clients (firm_id,name,pan,contact_person,contact_phone,"
                           "contact_email,address,entity_type,created_by_id) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)")
                    if qry_id:
                        client_id = qry_id(sql, args)
                    else:
                        execute(sql, args)
                        row_c = qry("SELECT id FROM clients WHERE firm_id=%s AND name=%s ORDER BY id DESC LIMIT 1",
                                    (g.firm_id, name), one=True)
                        client_id = row_c["id"] if row_c else None
                    clients_created += 1
                    if pan:
                        pan_index[pan.upper()] = {"id": client_id, "name": name, "pan": pan}
                    name_index[name.lower()] = {"id": client_id, "name": name, "pan": pan}
                except Exception as e:
                    row_errors.append(f"Row {rn}: could not create client — {e}")
                    continue
            if not client_id:
                continue

            packs_text = cell(row, packs_col) if packs_col else None
            packs_list = _parse_packs_cell(packs_text) if packs_text else list(default_packs)
            if packs_list:
                n = _apply_packs(client_id, name, packs_list, import_fy)
                packs_applied += 1
                calendar_added += n

            eng_cell = cell(row, "Type of Engagement")
            if not eng_cell:
                continue
            for t in _parse_engagement_types(eng_cell):
                display_type = t.split(":", 1)[1] if t.startswith("Other:") else t
                canonical_type = "Other" if t.startswith("Other:") else t
                k = eng_key(client_id, canonical_type, import_fy)
                fys_touched.add(import_fy)
                if k in eng_index:
                    engagements_skipped += 1
                    continue
                title = f"{name} – {display_type} FY {import_fy}"
                try:
                    esql = ("INSERT INTO engagements (firm_id,client_id,title,engagement_type,financial_year,created_by_id) "
                            "VALUES (%s,%s,%s,%s,%s,%s)")
                    eargs = (g.firm_id, client_id, title, canonical_type, import_fy, g.user["id"])
                    if qry_id:
                        qry_id(esql, eargs)
                    else:
                        execute(esql, eargs)
                    eng_index.add(k)
                    engagements_created += 1
                except Exception as e:
                    row_errors.append(f"Row {rn}: engagement failed — {e}")

        try:
            log_action(g.firm_id, g.user["id"], "IMPORT_CLIENTS", "Client", None,
                       f"Imported: {clients_created} clients, {engagements_created} engs, "
                       f"{packs_applied} packs, {calendar_added} cal (FY {import_fy})",
                       request.remote_addr)
        except Exception:
            pass

        return jsonify({
            "clients_created": clients_created, "clients_skipped": clients_skipped,
            "engagements_created": engagements_created, "engagements_skipped": engagements_skipped,
            "packs_applied": packs_applied, "calendar_items_added": calendar_added,
            "engagement_financial_years": sorted(fys_touched),
            "import_financial_year": import_fy,
            "errors": row_errors, "warnings": row_warnings, "rows_processed": len(df),
        }), 200

    try:
        @login_required
        def clients_import_template():
            return enhanced_template()

        @login_required
        def clients_import_excel():
            return enhanced_import()

        app.view_functions["clients_import_template"] = clients_import_template
        app.view_functions["clients_import_excel"] = clients_import_excel
        print("[import-packs] Clients Excel import supports Compliance Packs")
    except Exception as e:
        print("[import-packs] override failed:", e)
