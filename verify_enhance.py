"""
CA FirmHub — Data Verify (Inward vs Client vs Engagement).
Designed & Built by Harsh Pala.

Compares Doc Register (Inward) against clients and engagements for a FY,
and exports Excel for received vs pending analysis.
"""
from io import BytesIO
from datetime import datetime, date
from flask import request, jsonify, g, send_file


def register_verify_enhance(app, helpers):
    qry = helpers["qry"]
    login_required = helpers["login_required"]

    def _fy():
        return (request.args.get("fy") or "").strip() or None

    def _build_rows(firm_id, fy):
        eng_sql = (
            "SELECT e.id AS engagement_id, e.title AS engagement_title, "
            "e.engagement_type, e.financial_year, e.status AS eng_status, "
            "e.client_id, c.name AS client_name "
            "FROM engagements e "
            "JOIN clients c ON c.id = e.client_id "
            "WHERE e.firm_id=%s AND COALESCE(c.is_active, TRUE)=TRUE "
        )
        params = [firm_id]
        if fy:
            eng_sql += "AND e.financial_year=%s "
            params.append(fy)
        eng_sql += "ORDER BY c.name, e.title"
        engs = qry(eng_sql, tuple(params)) or []

        docs = qry(
            "SELECT dr.id, dr.client_id, dr.engagement_id, dr.doc_name, dr.doc_category, "
            "dr.doc_date, dr.status, dr.reference_no, dr.received_from "
            "FROM doc_register dr "
            "WHERE dr.firm_id=%s AND dr.doc_type='Inward'",
            (firm_id,)) or []

        by_eng = {}
        by_client = {}
        for d in docs:
            eid = d.get("engagement_id")
            cid = d.get("client_id")
            if eid:
                by_eng.setdefault(eid, []).append(d)
            if cid:
                by_client.setdefault(cid, []).append(d)

        def summarize(doc_list):
            received_statuses = {"Received", "Acknowledged", "Dispatched"}
            pending_statuses = {"Pending", "Returned"}
            total = len(doc_list)
            recv = sum(1 for x in doc_list if (x.get("status") or "") in received_statuses)
            pend = sum(1 for x in doc_list if (x.get("status") or "") in pending_statuses)
            other = total - recv - pend
            dates = [x.get("doc_date") for x in doc_list if x.get("doc_date")]
            last = max(dates) if dates else None
            if isinstance(last, datetime):
                last = last.date()
            if total == 0:
                status = "No data"
            elif pend == 0 and other == 0 and recv > 0:
                status = "Received"
            elif recv == 0 and pend > 0:
                status = "Pending"
            elif recv > 0 and (pend > 0 or other > 0):
                status = "Partial"
            else:
                status = "Other"
            return {
                "inward_count": total,
                "received_count": recv,
                "pending_count": pend,
                "other_count": other,
                "last_inward_date": str(last) if last else "",
                "verify_status": status,
            }

        rows = []
        for e in engs:
            eid = e["engagement_id"]
            linked = list(by_eng.get(eid, []))
            if not linked:
                linked = [d for d in by_client.get(e["client_id"], []) if not d.get("engagement_id")]
            sm = summarize(linked)
            rows.append({
                "client_id": e["client_id"],
                "client_name": e["client_name"],
                "engagement_id": eid,
                "engagement_title": e.get("engagement_title") or "",
                "engagement_type": e.get("engagement_type") or "",
                "financial_year": e.get("financial_year") or "",
                "eng_status": e.get("eng_status") or "",
                **sm,
            })

        if fy:
            clients_with_eng = {r["client_id"] for r in rows}
            for cid, dlist in by_client.items():
                if cid in clients_with_eng:
                    continue
                c = qry("SELECT name FROM clients WHERE id=%s AND firm_id=%s", (cid, firm_id), one=True)
                if not c:
                    continue
                sm = summarize(dlist)
                if sm["inward_count"] == 0:
                    continue
                rows.append({
                    "client_id": cid,
                    "client_name": c.get("name") or "",
                    "engagement_id": None,
                    "engagement_title": "(No engagement in FY)",
                    "engagement_type": "",
                    "financial_year": fy or "",
                    "eng_status": "",
                    **sm,
                })

        totals = {
            "engagements": len([r for r in rows if r.get("engagement_id")]),
            "clients": len({r["client_id"] for r in rows}),
            "inward_total": sum(r["inward_count"] for r in rows),
            "received_total": sum(r["received_count"] for r in rows),
            "pending_total": sum(r["pending_count"] for r in rows),
            "status_received": sum(1 for r in rows if r["verify_status"] == "Received"),
            "status_partial": sum(1 for r in rows if r["verify_status"] == "Partial"),
            "status_pending": sum(1 for r in rows if r["verify_status"] == "Pending"),
            "status_none": sum(1 for r in rows if r["verify_status"] == "No data"),
        }
        return rows, totals

    @app.route("/api/ops/verify/summary")
    @login_required
    def verify_summary():
        if g.user.get("role") not in ("Admin", "Team Leader", "Member", "SuperAdmin"):
            return jsonify({"detail": "Staff only"}), 403
        fy = _fy()
        rows, totals = _build_rows(g.firm_id, fy)
        return jsonify({"fy": fy, "rows": rows, "totals": totals})

    @app.route("/api/ops/verify/export")
    @login_required
    def verify_export():
        if g.user.get("role") not in ("Admin", "Team Leader", "Member", "SuperAdmin"):
            return jsonify({"detail": "Staff only"}), 403
        fy = _fy()
        rows, totals = _build_rows(g.firm_id, fy)

        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Data Verify"

        headers = [
            "Client Name", "Engagement", "Engagement Type", "Financial Year",
            "Eng Status", "Inward Docs", "Received", "Pending", "Other",
            "Last Inward Date", "Verify Status",
        ]
        header_fill = PatternFill("solid", fgColor="003366")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        thin = Border(
            left=Side(style="thin", color="D1DAEA"),
            right=Side(style="thin", color="D1DAEA"),
            top=Side(style="thin", color="D1DAEA"),
            bottom=Side(style="thin", color="D1DAEA"),
        )
        status_fills = {
            "Received": PatternFill("solid", fgColor="C6F6D5"),
            "Partial": PatternFill("solid", fgColor="FEFCBF"),
            "Pending": PatternFill("solid", fgColor="FED7D7"),
            "No data": PatternFill("solid", fgColor="E2E8F0"),
            "Other": PatternFill("solid", fgColor="E2E8F0"),
        }

        for col, h in enumerate(headers, 1):
            cell = ws.cell(1, col, h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = thin

        for i, r in enumerate(rows, 2):
            vals = [
                r.get("client_name"),
                r.get("engagement_title"),
                r.get("engagement_type"),
                r.get("financial_year"),
                r.get("eng_status"),
                r.get("inward_count", 0),
                r.get("received_count", 0),
                r.get("pending_count", 0),
                r.get("other_count", 0),
                r.get("last_inward_date") or "",
                r.get("verify_status"),
            ]
            for col, v in enumerate(vals, 1):
                cell = ws.cell(i, col, v)
                cell.border = thin
                if col == 11:
                    cell.fill = status_fills.get(str(v), status_fills["Other"])
                    cell.alignment = Alignment(horizontal="center")

        ws2 = wb.create_sheet("Summary")
        ws2["A1"] = "CA FirmHub — Data Verify Summary"
        ws2["A1"].font = Font(bold=True, size=14, color="003366")
        ws2["A2"] = f"Financial Year: {fy or 'All'}"
        ws2["A3"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        ws2["A4"] = "Designed & Built by Harsh Pala"
        labels = [
            ("Engagements", totals["engagements"]),
            ("Clients", totals["clients"]),
            ("Total inward docs", totals["inward_total"]),
            ("Received count", totals["received_total"]),
            ("Pending count", totals["pending_total"]),
            ("Engagements fully received", totals["status_received"]),
            ("Engagements partial", totals["status_partial"]),
            ("Engagements pending", totals["status_pending"]),
            ("Engagements with no data", totals["status_none"]),
        ]
        ws2["A6"] = "Metric"
        ws2["B6"] = "Value"
        ws2["A6"].font = Font(bold=True)
        ws2["B6"].font = Font(bold=True)
        for idx, (lab, val) in enumerate(labels, 7):
            ws2.cell(idx, 1, lab)
            ws2.cell(idx, 2, val)

        widths = [28, 28, 18, 14, 12, 12, 10, 10, 8, 14, 14]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws2.column_dimensions["A"].width = 32
        ws2.column_dimensions["B"].width = 14

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        fname = f"CA_FirmHub_Data_Verify_{fy or 'all'}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=fname,
        )

    print("[verify] Data verify routes registered")
